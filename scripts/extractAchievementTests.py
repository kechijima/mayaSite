#!/usr/bin/env python3
"""docs/到達度診断テスト/{赤系,白系,青系,黄系}.xlsx → scripts/achievementTests.data.ts

到達度診断テスト(太陽の紋章ごとの自己診断アンケート)の問題を、xlsx から seed 用の
TypeScript データに起こす。実行: python3 scripts/extractAchievementTests.py

元データの形(各ファイルに「★紋章名」シートが5枚、20紋章ぶん):
  A1        紋章名(コピーのままの誤記が多いので使わない。シート名が正)
  E2:G2     そう思う / どちらでもない / 思わない
  A列       カテゴリ名(思考 / 行動 / 人間関係 / 信念 / スキル。5行結合)
  B列       カテゴリの説明文(スキルだけ無い)
  C列 D列   問番号(1〜5)と問題文
  E,F,G列   その回答を選んだときの点数。[4,2,0] か [0,2,4] のどちらか
「★」の付かないシート(赤系の Sheet1・「赤い竜 (2)」、黄系の Sheet9)は作業途中のメモなので読まない。

データの手直し(2026-09-29、「もっともらしい方を採用」の方針で。CLAUDE.md にも記載):
  - 白い犬: スキル3「勘が鋭い方だ。（確認します）」、スキル5「…（角印します）」の注記を外す。
    末尾の「PDFがおかしい」はメモ行なので読まない。
  - 黄色い星: 信念5「昔から面倒見が良い方だ。→ プライドが高い方だと思う。（変更します）」は
    矢印の右側(変更後)「プライドが高い方だと思う。」を採用する。配点はそのまま [4,2,0]。
  - 黄色い太陽: 人間関係5の注記「（赤い竜と違うところです。）」を外す。
  - 文末の「。」の有無がまちまちなので、すべて「。」で終わらせる。
  - スキルの説明文は元データに無いので、他の4つと同じ調子でこちらで補った(SKILL_DESCRIPTION)。
"""
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / 'docs' / '到達度診断テスト'
OUT = ROOT / 'scripts' / 'achievementTests.data.ts'

# utils/mayaData.ts の SEALS と同じ順(sealIndex 0〜19)
SEAL_NAMES = [
    '赤い竜', '白い風', '青い夜', '黄色い種', '赤い蛇', '白い世界の橋渡し', '青い手', '黄色い星',
    '赤い月', '白い犬', '青い猿', '黄色い人', '赤い空歩く人', '白い魔法使い', '青い鷲', '黄色い戦士',
    '赤い地球', '白い鏡', '青い嵐', '黄色い太陽',
]
FILES = ['赤系.xlsx', '白系.xlsx', '青系.xlsx', '黄系.xlsx']
CATEGORY_KEYS = {'思考': 'thinking', '行動': 'action', '人間関係': 'relationship', '信念': 'belief', 'スキル': 'skill'}
SKILL_DESCRIPTION = 'スキルとは、思考・行動・人間関係・信念を日々の場面で発揮するために身についている力についての質問です。'

# 注記の除去と「→ 変更後」の採用
NOTE_RE = re.compile(r'[（(][^）)]*(確認します|角印します|変更します|違うところです)[^）)]*[）)]')


def clean_text(text: str) -> str:
    t = text.strip()
    t = NOTE_RE.sub('', t).strip()
    if '→' in t:
        t = t.split('→')[-1].strip()
    t = t.replace('　', ' ').strip()
    if not t.endswith('。'):
        t += '。'
    return t


def read_sheet(ws):
    cats = []
    cur = None
    for r in range(3, ws.max_row + 1):
        a, b, c, d, e, f, g = [ws.cell(r, col).value for col in range(1, 8)]
        if a:
            cur = {'name': a, 'description': (b or '').strip(), 'questions': []}
            cats.append(cur)
        if d is None or c is None:
            continue  # メモ行(白い犬の「PDFがおかしい」など)
        if not all(isinstance(v, (int, float)) for v in (e, f, g)):
            raise ValueError(f'{ws.title} 行{r}: 配点が数値ではない {e},{f},{g}')
        scores = [int(e), int(f), int(g)]
        if scores not in ([4, 2, 0], [0, 2, 4]):
            raise ValueError(f'{ws.title} 行{r}: 想定外の配点 {scores}')
        cur['questions'].append({'text': clean_text(str(d)), 'scores': scores})
    assert [c['name'] for c in cats] == list(CATEGORY_KEYS), f'{ws.title}: カテゴリ {[c["name"] for c in cats]}'
    for c in cats:
        assert len(c['questions']) == 5, f'{ws.title} {c["name"]}: {len(c["questions"])}問'
        if c['name'] == 'スキル' and not c['description']:
            c['description'] = SKILL_DESCRIPTION
    return cats


def ts_str(s: str) -> str:
    return "'" + s.replace('\\', '\\\\').replace("'", "\\'") + "'"


def main():
    seals = {}
    for f in FILES:
        wb = openpyxl.load_workbook(SRC / f, data_only=True)
        for ws in wb.worksheets:
            if not ws.title.startswith('★'):
                continue
            name = ws.title[1:].strip()
            if name not in SEAL_NAMES:
                raise ValueError(f'未知の紋章シート: {ws.title}')
            seals[name] = read_sheet(ws)
    missing = [n for n in SEAL_NAMES if n not in seals]
    assert not missing, f'足りない紋章: {missing}'

    lines = [
        '// scripts/extractAchievementTests.py が docs/到達度診断テスト/*.xlsx から生成したファイル。手で編集しない。',
        '// 元データの形と、取り込み時に直した箇所はそのスクリプトの冒頭コメントを参照。',
        '// 形式は utils/achievementTest.ts の AchievementTestDoc と対(scores は [そう思う, どちらでもない, 思わない])。',
        '',
        'export interface SeedQuestion {',
        '  text: string',
        '  scores: [number, number, number]',
        '}',
        '',
        'export interface SeedCategory {',
        "  key: 'thinking' | 'action' | 'relationship' | 'belief' | 'skill'",
        '  name: string',
        '  description: string',
        '  questions: SeedQuestion[]',
        '}',
        '',
        'export interface SeedTest {',
        '  sealIndex: number',
        '  sealName: string',
        '  categories: SeedCategory[]',
        '}',
        '',
        'export const ACHIEVEMENT_TEST_SEED: SeedTest[] = [',
    ]
    for i, name in enumerate(SEAL_NAMES):
        lines.append('  {')
        lines.append(f'    sealIndex: {i},')
        lines.append(f'    sealName: {ts_str(name)},')
        lines.append('    categories: [')
        for c in seals[name]:
            lines.append('      {')
            lines.append(f"        key: '{CATEGORY_KEYS[c['name']]}',")
            lines.append(f'        name: {ts_str(c["name"])},')
            lines.append(f'        description: {ts_str(c["description"])},')
            lines.append('        questions: [')
            for q in c['questions']:
                lines.append(f'          {{ text: {ts_str(q["text"])}, scores: [{", ".join(map(str, q["scores"]))}] }},')
            lines.append('        ]')
            lines.append('      },')
        lines.append('    ]')
        lines.append('  },')
    lines.append(']')
    lines.append('')
    OUT.write_text('\n'.join(lines), encoding='utf-8')
    n_q = sum(len(c['questions']) for s in seals.values() for c in s)
    print(f'wrote {OUT.relative_to(ROOT)}: {len(seals)} seals, {n_q} questions')


if __name__ == '__main__':
    main()
